#!/usr/bin/env python3
"""Собирает Excel-версию анализа из data/ideas.json и README.md.

Запуск:
    python3 scripts/build_xlsx.py                          # пишет ideas-russia-2026.xlsx в корень репозитория
    python3 scripts/build_xlsx.py --recalc=<recalc.py>     # то же + пересчёт в LibreOffice, чтобы в файле были значения

Балл и место в файле считаются формулами Excel от оценок и весов, поэтому таблицу можно
править прямо в Excel. Тексты карточек, кластеры, занятые ниши, правила рынка и источники
берутся из README.md, чтобы файл не расходился с отчётом.
"""
import json
import math
import re
import shutil
import tempfile
import zipfile
from fractions import Fraction
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data" / "ideas.json"
README = ROOT / "README.md"
OUT = ROOT / "ideas-russia-2026.xlsx"

FONT = "Arial"
TEAL = "0B6B62"
F_BODY = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK_SHEET = Font(name=FONT, size=10, color="008000")
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color="555555")
F_URL = Font(name=FONT, size=10, color="0563C1", underline="single")
FILL_HEAD = PatternFill("solid", fgColor=TEAL)
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="D0D7D4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER_TOP = Alignment(horizontal="center", vertical="top", wrap_text=True)
HEAD_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")

TYPE_LABEL = {"world": "Есть в мире", "new": "Новое"}
COMPETITION = [("Нет", 1.0), ("Низкая", 0.75), ("Средняя", 0.5), ("Высокая", 0.25)]
RISK = [("Низкий", 1.0), ("Средний", 0.6), ("Высокий", 0.2)]
WEIGHT_ROWS = [
    ("profit", "Прибыльность", "Потенциальная выручка через ~2 года при успехе"),
    ("popularity", "Популярность", "Потенциальный охват пользователей"),
    ("competition", "Свобода от конкуренции", "Насколько ниша в РФ ещё не занята"),
    ("ease", "Простота", "Обратная величина сложности реализации"),
    ("cheapness", "Дешевизна", "Обратная величина стоимости MVP"),
    ("reg_safety", "Регуляторная безопасность", "Обратная величина регуляторного риска"),
]
CARD_FIELDS = ["Что это", "Зачем", "Почему взлетит в России", "Конкуренты", "Как сделать", "Как зарабатывать", "Риски"]

LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def strip_md(text):
    """Убирает markdown-разметку и возвращает (чистый текст, список внешних ссылок)."""
    urls = [u for _, u in LINK_RE.findall(text) if u.startswith("http")]
    text = LINK_RE.sub(r"\1", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"(?<![\*\w])\*(?!\s)(.+?)(?<!\s)\*(?![\*\w])", r"\1", text)
    text = text.replace("`", "").replace("<br>", "\n")
    return text.strip(), urls


def section(readme, title):
    start = readme.index("\n## " + title)
    end = readme.find("\n## ", start + 4)
    return readme[start:end if end != -1 else len(readme)]


def md_table(block, first_header):
    lines = block[block.index("| " + first_header):].split("\n")
    rows = []
    for line in lines:
        if not line.startswith("|"):
            break
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if set("".join(cells)) <= set("-: "):
            continue
        rows.append(cells)
    return rows[0], rows[1:]


def parse_cards(readme):
    cards = {}
    for block in readme.split('\n<a id="')[1:]:
        head, _, body = block.partition("\n")
        cid = head.split('"')[0].upper()
        fields, current = {}, None
        subtitle = body.split("\n")[1].strip().strip("*").split(" · Балл")[0]
        fields["Аналоги (подробно)"] = re.sub(r"^Аналоги?: ", "", subtitle)
        for line in body.split("\n")[1:]:
            if line.startswith(("### ", "## ", "---", "<a id=")):
                break
            m = re.match(r"^- \*\*(.+?)\.\*\*\s*(.*)$", line)
            if m:
                current = m.group(1)
                if current not in CARD_FIELDS:
                    raise ValueError(f"{cid}: неизвестное поле карточки «{current}»")
                fields[current] = m.group(2)
            elif line.startswith("  - ") and current:
                fields[current] += "\n• " + line[4:]
            elif line.strip() and current and line.startswith("  "):
                fields[current] += " " + line.strip()
        cards[cid] = fields
    return cards


def parse_rules(readme):
    rows, topic = [], None
    for line in section(readme, "Правила игры в России в 2026 году").split("\n"):
        m = re.match(r"^\*\*(\d+)\. (.+?)\*\*\s*(.*)$", line)
        if m:
            topic = m.group(2).rstrip(".")
            if m.group(3):
                rows.append((topic, m.group(3)))
        elif line.startswith("- ") and topic:
            rows.append((topic, line[2:]))
    return rows


def parse_sources(readme):
    rows = []
    for line in section(readme, "Источники").split("\n"):
        m = re.match(r"^\*\*(.+?):\*\*(.*)$", line)
        if m:
            rows += [(m.group(1), t, u) for t, u in LINK_RE.findall(m.group(2))]
    return rows


def parse_bullets(readme, title):
    return [line[2:] for line in section(readme, title).split("\n") if line.startswith("- ")]


def score_formula(r):
    w = "'Веса и шкалы'!"
    comp = f"INDEX({w}$B$16:$B$19,MATCH(N{r},{w}$A$16:$A$19,0))"
    risk = f"INDEX({w}$B$23:$B$25,MATCH(O{r},{w}$A$23:$A$25,0))"
    total = f"SUM({w}$C$6:$C$11)"
    body = (
        f"{w}$C$6*L{r}/5+{w}$C$7*M{r}/5+{w}$C$8*{comp}"
        f"+{w}$C$9*(6-H{r})/5+{w}$C$10*(4-I{r})/3+{w}$C$11*{risk}"
    )
    return f"=IF({total}=0,0,ROUND(100*({body})/{total},0))"


def landscape_fit(ws):
    """Печать: A4, альбомная ориентация, по ширине страницы."""
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def set_widths(ws, widths):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def fit_row_height(ws, row, widths, min_lines=1):
    """Оценивает высоту строки по длине текста в переносимых ячейках (Excel не подгоняет её сам).

    Оценка с запасом: в строку помещается ~0,9 ширины столбца обычным шрифтом и ~0,8 — жирным,
    для мелкого шрифта пропорционально больше.
    """
    height = min_lines * 12.75
    for col, width in widths.items():
        cell = ws[f"{col}{row}"]
        value = cell.value
        if not isinstance(value, str) or value.startswith("="):
            continue
        size = cell.font.sz or 10
        per_line = max(1.0, width * (0.8 if cell.font.b else 0.9) * (10 / size))
        line_h = 11.25 if size <= 8 else 12.75
        n = sum(max(1, math.ceil(len(part) / per_line)) for part in value.split("\n"))
        height = max(height, n * line_h)
    ws.row_dimensions[row].height = min(409, height + 4)


def header_row(ws, row, headers):
    for col, text in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=text)
        c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD, HEAD_ALIGN, BORDER


def body_cell(ws, row, col, value, font=F_BODY, align=WRAP_TOP):
    c = ws.cell(row=row, column=col, value=value)
    c.font, c.alignment, c.border = font, align, BORDER
    return c


def weight_ints(meta):
    """Переводит веса из data/ideas.json в целые числа 0–10 для ячеек листа «Веса и шкалы»."""
    ints = {k: round(v * 20) for k, v in meta["weights"].items()}
    total_w, total_i = sum(meta["weights"].values()), sum(ints.values())
    for k, v in meta["weights"].items():
        if not 0 <= ints[k] <= 10 or abs(ints[k] / total_i - v / total_w) > 1e-9:
            raise SystemExit(f"Вес {k}={v} не переводится в целое число 0–10 без искажения пропорций")
    return ints


def build_weights(wb, meta, readme):
    ws = wb.create_sheet("Веса и шкалы")
    ws["A1"] = "Веса итогового балла и шкалы оценок"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Меняйте жёлтые ячейки (синий шрифт) — балл и место на листе «Сводная» пересчитаются. "
                "Значения по умолчанию подобраны под небольшую команду с ограниченным бюджетом.")
    ws["A2"].font = F_NOTE
    ws["A2"].alignment = WRAP_TOP
    ws.merge_cells("A2:D2")
    ws.row_dimensions[2].height = 28

    header_row(ws, 5, ["Критерий", "Что оценивает", "Вес (ввод)", "Доля в балле"])
    defaults = weight_ints(meta)
    for n, (key, name, desc) in enumerate(WEIGHT_ROWS):
        r = 6 + n
        body_cell(ws, r, 1, name, F_BOLD)
        body_cell(ws, r, 2, desc)
        c = body_cell(ws, r, 3, defaults[key], F_INPUT, CENTER_TOP)
        c.fill = FILL_INPUT
        c = body_cell(ws, r, 4, f"=IF(SUM($C$6:$C$11)=0,0,C{r}/SUM($C$6:$C$11))", F_BODY, CENTER_TOP)
        c.number_format = "0%"
    body_cell(ws, 12, 1, "Итого", F_BOLD)
    body_cell(ws, 12, 3, "=SUM(C6:C11)", F_BOLD, CENTER_TOP)
    c = body_cell(ws, 12, 4, "=SUM(D6:D11)", F_BOLD, CENTER_TOP)
    c.number_format = "0%"
    ws.add_data_validation(dv := DataValidation(type="whole", operator="between", formula1="0", formula2="10",
                                                 showErrorMessage=True, error="Введите целое число от 0 до 10"))
    dv.add("C6:C11")

    ws["A14"] = "Перевод текстовых оценок в доли (0–1)"
    ws["A14"].font = F_BOLD
    header_row(ws, 15, ["Конкуренция в РФ", "Свобода (0–1)"])
    for n, (label, value) in enumerate(COMPETITION):
        body_cell(ws, 16 + n, 1, label)
        c = body_cell(ws, 16 + n, 2, value, F_INPUT, CENTER_TOP)
        c.fill = FILL_INPUT
    header_row(ws, 22, ["Регуляторный риск", "Безопасность (0–1)"])
    for n, (label, value) in enumerate(RISK):
        body_cell(ws, 23 + n, 1, label)
        c = body_cell(ws, 23 + n, 2, value, F_INPUT, CENTER_TOP)
        c.fill = FILL_INPUT
    share = DataValidation(type="decimal", operator="between", formula1="0", formula2="1",
                           showErrorMessage=True, error="Введите долю от 0 до 1 (например, 0,75)")
    ws.add_data_validation(share)
    share.add("B16:B19")
    share.add("B23:B25")

    ws.row_breaks.append(Break(id=27))  # «Формула балла» не отрывается от своего текста при печати
    ws["A28"] = "Формула балла"
    ws["A28"].font = F_BOLD
    ws["A29"] = ("Балл = ОКРУГЛ(100 × (Вес прибыльности × Прибыльность/5 + Вес популярности × Популярность/5 "
                 "+ Вес конкуренции × Свобода + Вес простоты × (6 − Сложность)/5 + Вес дешевизны × (4 − Уровень стоимости)/3 "
                 "+ Вес безопасности × Безопасность) / Сумма весов). Место — по убыванию балла, при равенстве — по ID.")
    ws["A29"].font = F_BODY
    ws["A29"].alignment = WRAP_TOP
    ws.merge_cells("A29:D29")
    ws.row_dimensions[29].height = 54

    ws["A31"] = "Шкалы оценок"
    ws["A31"].font = F_BOLD
    header_row(ws, 32, ["Шкала", "Значения"])
    ws.merge_cells("B32:D32")
    scale_names = {"complexity": "Сложность (1–5)", "cost_tier": "Стоимость MVP (уровень 1–4)",
                   "profit": "Прибыльность (1–5)", "popularity": "Популярность (1–5)",
                   "competition": "Конкуренция в РФ", "reg_risk": "Регуляторный риск"}
    r = 33
    for key, name in scale_names.items():
        body_cell(ws, r, 1, name, F_BOLD)
        body_cell(ws, r, 2, meta["scales"][key])
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.row_dimensions[r].height = 40
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Методология и ограничения").font = F_BOLD
    for bullet in parse_bullets(readme, "Методология и ограничения"):
        r += 1
        text, _ = strip_md(bullet)
        c = ws.cell(row=r, column=1, value="• " + text)
        c.font, c.alignment = F_BODY, WRAP_TOP
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        ws.row_dimensions[r].height = max(28, math.ceil(len(text) / 120) * 12.75 + 4)

    set_widths(ws, {"A": 30, "B": 52, "C": 14, "D": 30})
    landscape_fit(ws)


def build_summary(wb, ideas):
    ws = wb.active
    ws.title = "Сводная"
    ws["A1"] = "Онлайн-сервисы для России: 34 идеи, которые можно сделать с Claude Code (срез на 20.10.2026)"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Синие ячейки и цветные столбцы «Тип», «Конкуренция в РФ», «Регуляторный риск» — экспертные оценки, "
                "их можно менять (в цветных столбцах — выпадающие списки).")
    ws["A3"] = ("Балл, место и стоимость в рублях считаются формулами; веса — на листе «Веса и шкалы». "
                "Подробное описание каждой идеи — на листе «Карточки».")
    ws["A2"].font = ws["A3"].font = F_NOTE

    headers = ["Место", "ID", "Продукт", "Тип", "Балл\n(0–100)", "Категория", "Аналог в мире",
               "Сложность (1–5)", "Стоимость MVP, уровень (1–4)", "Стоимость MVP", "Срок MVP",
               "Прибыльность (1–5)", "Популярность (1–5)", "Конкуренция в РФ", "Регуляторный риск",
               "Монетизация", "Канал запуска"]
    header_row(ws, 4, headers)
    ws.row_dimensions[4].height = 42
    notes = {
        "A": "Место по баллу (формула). При равном балле выше идёт меньший ID.",
        "D": "«Есть в мире» — сервис работает за рубежом, а в РФ его нет или он в зачатке. «Новое» — массовых аналогов не найдено нигде.",
        "E": "Взвешенная сумма шести оценок, считается формулой. Веса — на листе «Веса и шкалы».",
        "H": "Экспертная оценка. 1 — MVP за 1–3 недели одним разработчиком с Claude Code; 2 — 1–2 месяца; 3 — 2–4 месяца; 4 — 4–6+ месяцев; 5 — больше 8 месяцев.",
        "I": "Экспертная оценка. 1 — до ~0,3 млн ₽; 2 — 0,3–1 млн; 3 — 1–3 млн; 4 — больше 3 млн. Без зарплаты основателя и маркетинга.",
        "L": "Экспертная оценка выручки в месяц через ~2 года при успехе: 1 — до 0,5 млн ₽; 2 — 0,5–2 млн; 3 — 2–10 млн; 4 — 10–50 млн; 5 — больше 50 млн.",
        "M": "Экспертная оценка охвата: 1 — до 50 тыс. пользователей; 2 — до 300 тыс.; 3 — до 1–2 млн; 4 — до 5–10 млн; 5 — больше 10 млн.",
        "N": "Насколько ниша в РФ уже занята, по результатам поиска на октябрь 2026 года.",
        "O": "Регуляторные и правовые риски в РФ: персональные данные, медицина, телефония, платформы.",
    }
    for col, text in notes.items():
        ws[f"{col}4"].comment = Comment(text, "Анализ")

    first, last = 5, 4 + len(ideas)
    for n, i in enumerate(ideas):
        r = first + n
        body_cell(ws, r, 1, f"=RANK(E{r},$E${first}:$E${last},0)+COUNTIFS($E${first}:$E${last},E{r},$B${first}:$B${last},\"<\"&B{r})",
                  F_BOLD, CENTER_TOP)
        body_cell(ws, r, 2, i["id"], F_BODY, CENTER_TOP)
        body_cell(ws, r, 3, i["name"], F_BOLD)
        body_cell(ws, r, 4, TYPE_LABEL[i["type"]], F_INPUT)
        body_cell(ws, r, 5, score_formula(r), Font(name=FONT, size=11, bold=True), CENTER_TOP)
        body_cell(ws, r, 6, i["category"])
        body_cell(ws, r, 7, i["analog"])
        body_cell(ws, r, 8, i["complexity"], F_INPUT, CENTER_TOP)
        body_cell(ws, r, 9, i["cost_tier"], F_INPUT, CENTER_TOP)
        body_cell(ws, r, 10, f'=CHOOSE(I{r},"до 0,3 млн ₽","0,3–1 млн ₽","1–3 млн ₽","3+ млн ₽")', F_BODY, CENTER_TOP)
        body_cell(ws, r, 11, i["mvp_time"], F_BODY, CENTER_TOP)
        body_cell(ws, r, 12, i["profit"], F_INPUT, CENTER_TOP)
        body_cell(ws, r, 13, i["popularity"], F_INPUT, CENTER_TOP)
        body_cell(ws, r, 14, i["competition"], F_INPUT, CENTER_TOP)
        body_cell(ws, r, 15, i["reg_risk"], F_INPUT, CENTER_TOP)
        body_cell(ws, r, 16, i["monetization"])
        body_cell(ws, r, 17, i["channel"])

    widths = {"A": 8, "B": 6, "C": 44, "D": 13, "E": 10, "F": 14, "G": 28, "H": 11, "I": 14, "J": 13,
              "K": 11, "L": 15, "M": 15, "N": 14, "O": 15, "P": 34, "Q": 30}
    set_widths(ws, widths)
    for r in range(first, last + 1):
        fit_row_height(ws, r, {k: widths[k] for k in "CGPQ"})

    rng = lambda col: f"{col}{first}:{col}{last}"
    validations = [
        (DataValidation(type="list", formula1='"Есть в мире,Новое"'), "D"),
        (DataValidation(type="whole", operator="between", formula1="1", formula2="5"), "H"),
        (DataValidation(type="whole", operator="between", formula1="1", formula2="4"), "I"),
        (DataValidation(type="whole", operator="between", formula1="1", formula2="5"), "L"),
        (DataValidation(type="whole", operator="between", formula1="1", formula2="5"), "M"),
        (DataValidation(type="list", formula1='"Нет,Низкая,Средняя,Высокая"'), "N"),
        (DataValidation(type="list", formula1='"Низкий,Средний,Высокий"'), "O"),
    ]
    for dv, col in validations:
        dv.showErrorMessage = True
        dv.error = "Выберите значение из списка или шкалы"
        ws.add_data_validation(dv)
        dv.add(rng(col))

    cf = ws.conditional_formatting
    cf.add(rng("E"), ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="63BE7B"))
    for col in "LM":
        cf.add(rng(col), ColorScaleRule(start_type="num", start_value=1, start_color="FFFFFF",
                                         end_type="num", end_value=5, end_color="63BE7B"))
    cf.add(rng("H"), ColorScaleRule(start_type="num", start_value=1, start_color="FFFFFF",
                                     end_type="num", end_value=5, end_color="F8A86B"))
    cf.add(rng("I"), ColorScaleRule(start_type="num", start_value=1, start_color="FFFFFF",
                                     end_type="num", end_value=4, end_color="F8A86B"))
    palette = {"good": ("E0F1E5", "1B743B"), "warn": ("F7ECD2", "8A5800"), "bad": ("FAE3E0", "AD241C")}

    def colour(col, value, tone):
        bg, fg = palette[tone]
        cf.add(rng(col), CellIsRule(operator="equal", formula=[f'"{value}"'],
                                    fill=PatternFill("solid", fgColor=bg, bgColor=bg),
                                    font=Font(name=FONT, size=10, bold=True, color=fg)))

    for value, tone in [("Нет", "good"), ("Низкая", "good"), ("Средняя", "warn"), ("Высокая", "bad")]:
        colour("N", value, tone)
    for value, tone in [("Низкий", "good"), ("Средний", "warn"), ("Высокий", "bad")]:
        colour("O", value, tone)
    cf.add(rng("D"), CellIsRule(operator="equal", formula=['"Новое"'],
                                fill=PatternFill("solid", fgColor="D9EEEA", bgColor="D9EEEA"),
                                font=Font(name=FONT, size=10, bold=True, color=TEAL)))

    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:Q{last}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 2
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:4"
    ws.print_title_cols = "A:C"
    return first, last


def build_cards(wb, ideas, cards, first, last):
    ws = wb.create_sheet("Карточки")
    headers = ["ID", "Продукт", "Тип", "Место", "Балл", "Аналог в мире", "Аналоги (подробно)"] + CARD_FIELDS + ["Ссылки"]
    header_row(ws, 1, headers)
    ws.row_dimensions[1].height = 30
    s = "'Сводная'!"
    for n, i in enumerate(sorted(ideas, key=lambda x: x["id"])):
        r = 2 + n
        fields = cards[i["id"]]
        body_cell(ws, r, 1, i["id"], F_BOLD, CENTER_TOP)
        body_cell(ws, r, 2, i["name"], F_BOLD)
        for col, src in ((3, "D"), (4, "A"), (5, "E")):
            body_cell(ws, r, col, f"=INDEX({s}${src}${first}:${src}${last},MATCH(A{r},{s}$B${first}:$B${last},0))",
                      F_LINK_SHEET, CENTER_TOP)
        body_cell(ws, r, 6, i["analog"])
        body_cell(ws, r, 7, strip_md(fields["Аналоги (подробно)"])[0])
        urls = []
        for k, field in enumerate(CARD_FIELDS):
            text, found = strip_md(fields.get(field, ""))
            urls += found
            body_cell(ws, r, 8 + k, text or "—")
        body_cell(ws, r, 8 + len(CARD_FIELDS), "\n".join(dict.fromkeys(urls)) or "—", Font(name=FONT, size=8, color="555555"))

    widths = {"A": 6, "B": 28, "C": 12, "D": 8, "E": 7, "F": 22, "G": 28, "H": 46, "I": 34, "J": 46, "K": 40,
              "L": 40, "M": 38, "N": 38, "O": 40}
    set_widths(ws, widths)
    for r in range(2, 2 + len(ideas)):
        fit_row_height(ws, r, {k: widths[k] for k in "BFGHIJKLMNO"})
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:O{1 + len(ideas)}"
    # Карточки слишком широкие для печати «по ширине», поэтому на каждой странице повторяем шапку и ID/Продукт.
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape"
    ws.print_title_rows = "1:1"
    ws.print_title_cols = "A:B"


def build_text_sheet(wb, title, headers, rows, widths, url_col=None):
    ws = wb.create_sheet(title)
    header_row(ws, 1, headers)
    ws.row_dimensions[1].height = 30
    for n, row in enumerate(rows):
        r = 2 + n
        for k, value in enumerate(row):
            body_cell(ws, r, 1 + k, value, F_BOLD if k == 0 else F_BODY)
        fit_row_height(ws, r, widths)
        if url_col is not None and row[url_col].startswith("http"):
            # Формула HYPERLINK, а не гиперссылка ячейки: такой текст переносится и в LibreOffice.
            url = row[url_col]
            if len(url) > 255 or '"' in url:
                raise SystemExit(f"Ссылку нельзя записать в HYPERLINK: {url}")
            c = ws.cell(row=r, column=1 + url_col, value=f'=HYPERLINK("{url}","{url}")')
            c.font = F_URL
    set_widths(ws, widths)
    ws.freeze_panes = "B2"
    landscape_fit(ws)
    ws.print_title_rows = "1:1"
    return ws


def dedupe_print_titles(xml):
    """LibreOffice пишет область печати с повтором строк и столбцов дважды — оставляем одно определение на лист."""
    seen = set()

    def keep_first(m):
        sheet = re.search(r'localSheetId="(\d+)"', m.group(0))
        key = sheet.group(1) if sheet else None
        if key in seen:
            return ""
        seen.add(key)
        return m.group(0)

    return re.sub(r'<definedName[^>]*name="_xlnm\.Print_Titles"[^>]*>.*?</definedName>', keep_first, xml)


def place_notes(vml):
    """LibreOffice ставит все примечания в одну точку — переносим каждое рядом со своей ячейкой."""
    def fix(shape):
        block = shape.group(0)
        col = re.search(r"<x:Column>(\d+)</x:Column>", block)
        anchor = re.search(r"<x:Anchor>(.*?)</x:Anchor>", block, re.S)
        if not col or not anchor:
            return block
        a = [v.strip() for v in anchor.group(1).split(",")]
        c = int(col.group(1))
        a[0], a[1], a[4], a[5] = str(c + 1), "15", str(c + 3), "15"
        return block.replace(anchor.group(0), "<x:Anchor>" + ", ".join(a) + "</x:Anchor>")

    return re.sub(r"<v:shape\b.*?</v:shape>", fix, vml, flags=re.S)


def recalc_and_fix(recalc_script):
    """Пересчитывает формулы в LibreOffice и возвращает автора комментариев, которого LibreOffice заменяет."""
    import subprocess
    import sys
    out = subprocess.run([sys.executable, recalc_script, str(OUT), "120"], capture_output=True, text=True)
    report = json.loads(out.stdout)
    if report.get("status") != "success" or report.get("total_errors"):
        raise SystemExit(f"Пересчёт не удался: {report}")
    with tempfile.TemporaryDirectory() as tmp:
        fixed = Path(tmp) / OUT.name
        with zipfile.ZipFile(OUT) as src, zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                payload = src.read(item.filename)
                if item.filename.startswith("xl/comments"):
                    payload = payload.replace(b"Unknown Author", "Анализ".encode("utf-8"))
                elif item.filename.startswith("xl/drawings/vmlDrawing"):
                    payload = place_notes(payload.decode("utf-8")).encode("utf-8")
                elif item.filename == "xl/workbook.xml":
                    payload = dedupe_print_titles(payload.decode("utf-8")).encode("utf-8")
                dst.writestr(item, payload)
        shutil.move(fixed, OUT)
    print(f"Пересчитано формул: {report['total_formulas']}, ошибок: 0")


def main():
    import sys
    data = json.loads(DATA.read_text(encoding="utf-8"))
    readme = README.read_text(encoding="utf-8")
    ideas = data["ideas"]

    # Порядок строк — по баллу при весах по умолчанию (тот же, что в README). Место пересчитывается формулой.
    w = weight_ints(data["meta"])
    comp = {k: Fraction(str(v)) for k, v in COMPETITION}
    risk = {k: Fraction(str(v)) for k, v in RISK}

    def default_score(i):
        parts = (w["profit"] * Fraction(i["profit"], 5) + w["popularity"] * Fraction(i["popularity"], 5)
                 + w["competition"] * comp[i["competition"]] + w["ease"] * Fraction(6 - i["complexity"], 5)
                 + w["cheapness"] * Fraction(4 - i["cost_tier"], 3) + w["reg_safety"] * risk[i["reg_risk"]])
        return math.floor(100 * parts / sum(w.values()) + Fraction(1, 2))  # как ROUND в Excel

    ideas.sort(key=lambda i: (-default_score(i), i["id"]))

    cards = parse_cards(readme)
    missing = {i["id"] for i in ideas} - set(cards)
    if missing:
        raise SystemExit(f"В README нет карточек: {sorted(missing)}")

    wb = Workbook()
    first, last = build_summary(wb, ideas)
    build_cards(wb, ideas, cards, first, last)
    build_weights(wb, data["meta"], readme)

    _, cluster_rows = md_table(section(readme, "С чего начать"), "Кластер")
    clusters = []
    for row in cluster_rows:
        cells = [strip_md(c)[0] for c in row]
        recommended = "⭐" in cells[0]
        clusters.append([cells[0].replace("⭐", "").replace("лучший", "").strip(), *cells[1:],
                         "Рекомендуем начать с него" if recommended else ""])
    build_text_sheet(wb, "Кластеры", ["Кластер", "Продукты (→ — порядок запуска)", "Кто платит", "Почему вместе", "Рекомендация"],
                     clusters, {"A": 24, "B": 50, "C": 22, "D": 60, "E": 22})

    _, niche_rows = md_table(section(readme, "Куда лучше не идти: ниши уже заняты"), "Ниша")
    niches = []
    for row in niche_rows:
        parsed = [strip_md(c) for c in row]
        urls = [u for _, found in parsed for u in found]
        niches.append([p[0] for p in parsed] + ["\n".join(urls) or "—"])
    build_text_sheet(wb, "Занятые ниши", ["Ниша (мировой аналог)", "Кто уже работает в России", "Вывод", "Ссылки"],
                     niches, {"A": 34, "B": 48, "C": 40, "D": 50})

    second = []
    for bullet in parse_bullets(readme, "Идеи второго эшелона (рассматривал, но не включил в таблицу)"):
        m = re.match(r"^\*\*(.+?)\*\*(.*)$", bullet)
        name, rest = (m.group(1), m.group(2)) if m else ("", bullet)
        if "**" in rest:  # пункт-перечисление из нескольких идей
            name, comment = "Ещё несколько идей", strip_md(bullet)[0]
        elif rest.startswith(":"):
            comment = strip_md(rest[1:])[0].strip()
            comment = comment[:1].upper() + comment[1:]
        else:
            head, _, tail = rest.partition(". ")
            sep = " " if head[:1].isspace() else ""
            name, comment = name + sep + strip_md(head)[0].rstrip("."), strip_md(tail)[0]
        second.append([name.strip(), comment.strip() or "—"])
    build_text_sheet(wb, "Второй эшелон", ["Идея", "Комментарий"], second, {"A": 46, "B": 90})

    rules = []
    for topic, text in parse_rules(readme):
        clean, urls = strip_md(text)
        rules.append([topic, clean, "\n".join(urls) or "—"])
    build_text_sheet(wb, "Правила рынка 2026", ["Тема", "Суть", "Ссылки"], rules, {"A": 30, "B": 90, "C": 50})

    build_text_sheet(wb, "Источники", ["Группа", "Источник", "Ссылка"], parse_sources(readme),
                     {"A": 24, "B": 50, "C": 80}, url_col=2)

    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.properties.title = "Онлайн-сервисы для России — 34 идеи (октябрь 2026)"
    wb.properties.creator = "Анализ"
    wb.properties.language = "ru-RU"
    wb.save(OUT)
    print(f"Сохранено: {OUT.relative_to(ROOT)}")
    recalc = [a.split("=", 1)[1] for a in sys.argv[1:] if a.startswith("--recalc=")]
    if recalc:
        recalc_and_fix(recalc[0])


if __name__ == "__main__":
    main()
